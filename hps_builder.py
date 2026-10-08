"""
hps_builder.py — Final HPS (RAB) Generator
==========================================

Mengambil data JSON hasil ekstraksi dari bot dan menyusunnya menjadi 
dokumen Excel (HPS / RAB) yang rapi.
"""

import json
import os
import re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

def generate_hps_excel(json_filepath: str, output_dir: str = None) -> str:
    """Membaca file JSON dan membuat file Excel HPS/RAB."""
    if output_dir is None:
        output_dir = os.path.dirname(json_filepath)
        
    with open(json_filepath, 'r', encoding='utf-8') as f:
        data = json.load(f)
        
    meta = data.get("metadata", {})
    kak = data.get("kak", {})
    hps = data.get("hps", {})
    
    judul = meta.get("judul_pekerjaan", "PEKERJAAN TIDAK DIKETAHUI").upper()
    tahun = meta.get("tahun_anggaran", datetime.now().year)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "HPS"
    
    # Define styles
    font_bold = Font(name='Arial', size=11, bold=True)
    font_normal = Font(name='Arial', size=11)
    
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='top')
    align_middle_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_middle_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # Set column widths
    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 2
    ws.column_dimensions['C'].width = 6
    ws.column_dimensions['D'].width = 6
    ws.column_dimensions['E'].width = 45
    ws.column_dimensions['F'].width = 8
    ws.column_dimensions['G'].width = 4
    ws.column_dimensions['H'].width = 8
    ws.column_dimensions['I'].width = 8
    ws.column_dimensions['J'].width = 18
    ws.column_dimensions['K'].width = 20
    ws.column_dimensions['L'].width = 2
    ws.column_dimensions['M'].width = 35
    
    # Formatting numbers
    total = hps.get("total", 0)
    total_str = f"{total:,.0f}".replace(",", ".")
    
    # Headers
    ws.merge_cells('C2:K2')
    ws['C2'] = "HARGA PERKIRAAN SENDIRI (HPS)"
    ws['C2'].font = font_bold
    ws['C2'].alignment = align_center
    
    ws.merge_cells('C3:K3')
    tipe_desc = "BELANJA JASA KONSULTANSI" if "konsultansi" in judul.lower() else "BELANJA BARANG / JASA"
    ws['C3'] = f"{tipe_desc}\n{judul}"
    ws['C3'].font = font_bold
    ws['C3'].alignment = align_center
    ws.row_dimensions[3].height = 40
    
    ws.merge_cells('C6:K6')
    ws['C6'] = f"TAHUN ANGGARAN {tahun}"
    ws['C6'].font = font_bold
    ws['C6'].alignment = align_center
    
    ws.merge_cells('C9:K9')
    ws['C9'] = f"ANGGARAN  : Rp. {total_str},-"
    ws['C9'].font = font_bold
    ws['C9'].alignment = align_left
    
    ws.merge_cells('C11:K11')
    ws['C11'] = "HARGA PERKIRAAN SENDIRI (HPS)"
    ws['C11'].font = font_bold
    ws['C11'].alignment = align_center
    
    # Table Header (Row 12)
    headers = {
        'C': "Nomor", 'D': "", 'E': "Rincian pekerjaan", 
        'F': "Volume", 'G': "", 'H': "", 'I': "", 
        'J': "Harga Satuan\n(Rp)", 'K': "Jumlah\n(Rp)", 'M': "Keterangan"
    }
    for col_letter, title in headers.items():
        cell = ws[f"{col_letter}12"]
        cell.value = title
        cell.font = font_bold
        cell.alignment = align_center
        # Applying borders for the table area (C to K)
        if col_letter in 'CDEFGHIJK':
            cell.border = thin_border
            
    # Merge Nomor and Rincian header if needed, but CSV didn't. 
    # CSV merged C12:D12 ? CSV has ",,Nomor,,Rincian pekerjaan" (so C is Nomor, D is empty).
    ws.merge_cells('C12:D12')
    ws.merge_cells('F12:I12') # Volume merges over F, G, H, I
    
    # Re-apply borders after merge
    for c in ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
        ws[f"{c}12"].border = thin_border
    
    current_row = 13
    
    def write_row(col_c="", col_d="", col_e="", col_f="", col_g="", col_h="", col_i="", col_j="", col_k="", col_m="", is_bold=False):
        nonlocal current_row
        ws[f'C{current_row}'] = col_c
        ws[f'D{current_row}'] = col_d
        ws[f'E{current_row}'] = col_e
        ws[f'F{current_row}'] = col_f
        ws[f'G{current_row}'] = col_g
        ws[f'H{current_row}'] = col_h
        ws[f'I{current_row}'] = col_i
        ws[f'J{current_row}'] = col_j
        ws[f'K{current_row}'] = col_k
        ws[f'M{current_row}'] = col_m
        
        ws[f'C{current_row}'].alignment = align_center
        ws[f'D{current_row}'].alignment = align_center
        ws[f'E{current_row}'].alignment = align_left
        ws[f'F{current_row}'].alignment = align_center
        ws[f'G{current_row}'].alignment = align_center
        ws[f'H{current_row}'].alignment = align_center
        ws[f'I{current_row}'].alignment = align_center
        ws[f'J{current_row}'].alignment = align_right
        ws[f'K{current_row}'].alignment = align_right
        ws[f'M{current_row}'].alignment = align_left
        
        if isinstance(col_j, (int, float)): ws[f'J{current_row}'].number_format = '#,##0.00'
        if isinstance(col_k, (int, float)): ws[f'K{current_row}'].number_format = '#,##0.00'
        
        font = font_bold if is_bold else font_normal
        for c in ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']:
            ws[f'{c}{current_row}'].border = thin_border
            ws[f'{c}{current_row}'].font = font
        ws[f'M{current_row}'].font = font
        
        current_row += 1

    biaya_personil = hps.get("biaya_personil", [])
    subtotal_personil = sum(item.get("jumlah", 0) for item in biaya_personil)
    write_row(col_c="I", col_e="Biaya Personil", col_k=subtotal_personil, is_bold=True)
    
    for i, item in enumerate(biaya_personil, 1):
        uraian = item.get("uraian", "")
        vol = item.get("volume", 1)
        sw = item.get("satuan_waktu", 1)
        sat = item.get("satuan", "")
        hs = item.get("harga_satuan", 0)
        jum = item.get("jumlah", vol * sw * hs)
        kualifikasi = item.get("kualifikasi", "")
        # Format: col_c="", col_d="1", col_e="Tenaga Ahli", col_f=1, col_g="x", col_h=3, col_i="ob"
        write_row(col_d=i, col_e=uraian, col_f=vol, col_g="x", col_h=sw, col_i=sat, col_j=hs, col_k=jum, col_m=kualifikasi)
        
    current_row += 1 # Empty row between sections
    
    biaya_non_personil = hps.get("biaya_non_personil", [])
    subtotal_non = sum(item.get("jumlah", 0) for item in biaya_non_personil)
    write_row(col_c="II", col_e="Biaya Non Personil", col_k=subtotal_non, is_bold=True)
    
    for i, item in enumerate(biaya_non_personil, 1):
        uraian = item.get("uraian", "")
        vol = item.get("volume", 1)
        sw = item.get("satuan_waktu", 1)
        sat = item.get("satuan", "")
        hs = item.get("harga_satuan", 0)
        jum = item.get("jumlah", vol * sw * hs)
        keterangan = item.get("keterangan", "")
        
        # Kadang non personil tidak ada satuan waktu (hanya 1)
        # Sesuai CSV: 1, x, 1, pcs. Kita akan tampilkan sw jika ada, atau 1.
        if sw in [None, 0, ""]: sw = 1
        
        write_row(col_d=i, col_e=uraian, col_f=vol, col_g="x", col_h=sw, col_i=sat, col_j=hs, col_k=jum, col_m=keterangan)

    current_row += 1 # Empty row
    
    subtotal = hps.get("subtotal", 0)
    if subtotal == 0: subtotal = subtotal_personil + subtotal_non
    
    write_row(col_e="Jumlah", col_k=subtotal, is_bold=True)
    
    ppn = hps.get("ppn", 0)
    if kak.get("sumber_pendanaan", {}).get("termasuk_ppn", True) and ppn == 0:
        ppn = int(subtotal * 0.11)
        total = subtotal + ppn
    
    write_row(col_e="PPN 11%", col_k=ppn, is_bold=True)
    write_row(col_e="JUMLAH TOTAL", col_k=total, is_bold=True)
    
    current_row += 2
    
    bulan_indo = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
    now = datetime.now()
    tanggal_ttd = f"Yogyakarta, {now.day} {bulan_indo[now.month-1]} {tahun}"
    
    # Signature Block
    for r, text in enumerate([tanggal_ttd, "PPKom,", "", "Raden Roro Andarini,S.E., M.Si.", "NIP. 19720317 199703 2 004"]):
        target_row = current_row + r
        ws.merge_cells(f'I{target_row}:K{target_row}')
        cell = ws[f'I{target_row}']
        cell.value = text
        cell.font = font_normal
        cell.alignment = align_center
    
    # Save file
    safe_title = "".join(c if c.isalnum() else "_" for c in judul)
    safe_title = re.sub(r"_+", "_", safe_title).strip("_")[:40]
    base_out_name = f"HPS_{safe_title}_{datetime.now().strftime('%d-%m-%Y_%H%M%S')}"
    
    if output_dir:
        folder_arsip = output_dir
    else:
        # Standarisasi Folder (Bisa pakai Arsip KAK)
        tahun_arsip = str(tahun)
        bulan_indo = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", 
                      "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
        nama_bulan = bulan_indo[datetime.now().month - 1]
        
        base_arsip = os.path.join(os.path.dirname(__file__), "Arsip KAK")
        folder_arsip = os.path.join(base_arsip, tahun_arsip, nama_bulan)
        os.makedirs(folder_arsip, exist_ok=True)
    
    filepath = os.path.join(folder_arsip, base_out_name + ".xlsx")
    wb.save(filepath)
    
    return filepath

if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        json_file = sys.argv[1]
        out_excel = generate_hps_excel(json_file)
        if out_excel:
            print(f"File Excel HPS berhasil dibuat: {out_excel}")
    else:
        print("Penggunaan: python hps_builder.py <file.json>")
